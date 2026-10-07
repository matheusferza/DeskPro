from __future__ import annotations

import json
import os
import re
from datetime import date, datetime, timedelta
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook, load_workbook

FIELD_ORDER = [
    "COD",
    "DATA ENTRADA",
    "CLIENTES",
    "PACIENTE",
    "SUP/INF",
    "SAIDA",
    "OBSERVAÇÕES",
    "ENTREGUE",
]

EXCEL_IMPORT_FIELDS = ["COD", "DATA ENTRADA", "CLIENTES", "PACIENTE", "SUP/INF", "SAIDA"]
EXCEL_EXPORT_FIELDS = ["COD", "DATA ENTRADA", "CLIENTES", "PACIENTE", "SUP/INF", "SAIDA", "STATUS", "OBSERVAÇÕES"]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def format_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    text = normalize_text(value)
    if not text:
        return ""
    try:
        parsed = parse_date(text)
        return parsed.strftime("%d/%m/%Y")
    except ValueError:
        return text


def parse_date(value: Any) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = normalize_text(value)
    if not text:
        return None

    today = date.today()

    match_day = re.fullmatch(r"\d{1,2}", text)
    if match_day:
        day = int(text)
        return date(today.year, today.month, day)

    match_day_month = re.fullmatch(r"\d{1,2}/\d{1,2}", text)
    if match_day_month:
        day_str, month_str = text.split("/")
        day = int(day_str)
        month = int(month_str)
        return date(today.year, month, day)

    match_day_month_year = re.fullmatch(r"\d{1,2}/\d{1,2}/\d{2,4}", text)
    if match_day_month_year:
        day_str, month_str, year_str = text.split("/")
        day = int(day_str)
        month = int(month_str)
        year = int(year_str)
        if year < 100:
            year += 2000 if year <= 69 else 1900
        return date(year, month, day)

    for fmt in (
        "%d/%m/%Y",
        "%d/%m/%y",
        "%d-%m-%Y",
        "%d-%m-%y",
        "%Y-%m-%d",
        "%d/%m/%Y %H:%M:%S",
        "%Y/%m/%d",
    ):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    raise ValueError(f"Data inválida: {value}")


def compute_status(record: Dict[str, Any]) -> str:
    if bool(record.get("ENTREGUE")):
        return "Entregue"

    saida = record.get("SAIDA")
    if isinstance(saida, date):
        saida_date = saida
    else:
        try:
            saida_date = parse_date(saida)
        except ValueError:
            saida_date = None

    if saida_date and saida_date < date.today():
        return "Atrasado"

    return "No prazo"


def normalize_sup_inf(value: Any) -> str:
    value_text = normalize_text(value)
    aliases = {
        "superior": "Superior",
        "inferior": "Inferior",
        "ambas": "Ambas",
        "sup": "Superior",
        "inf": "Inferior",
        "s": "Superior",
        "i": "Inferior",
        "a": "Ambas",
    }
    lower = value_text.lower()
    return aliases.get(lower, value_text or "")


def prepare_record(raw: Dict[str, Any]) -> Dict[str, Any]:
    record = {}
    for key in FIELD_ORDER:
        record[key] = raw.get(key, "")

    record["COD"] = normalize_text(record.get("COD"))
    record["DATA ENTRADA"] = format_date(record.get("DATA ENTRADA"))
    record["CLIENTES"] = normalize_text(record.get("CLIENTES"))
    record["PACIENTE"] = normalize_text(record.get("PACIENTE"))
    record["SUP/INF"] = normalize_sup_inf(record.get("SUP/INF"))
    record["SAIDA"] = format_date(record.get("SAIDA"))
    record["OBSERVAÇÕES"] = normalize_text(record.get("OBSERVAÇÕES"))

    entregue_value = record.get("ENTREGUE", False)
    if isinstance(entregue_value, str):
        record["ENTREGUE"] = entregue_value.strip().lower() in {"1", "true", "yes", "sim", "s"}
    else:
        record["ENTREGUE"] = bool(entregue_value)

    record["STATUS"] = compute_status(record)
    return record


def get_data_path() -> str:
    import sys
    if getattr(sys, "frozen", False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_dir, "labmarcos_data.json")


def load_records(storage_path: Optional[str] = None) -> List[Dict[str, Any]]:
    path = storage_path or get_data_path()
    if not os.path.exists(path):
        try:
            with open(path, "w", encoding="utf-8") as fh:
                json.dump([], fh, ensure_ascii=False, indent=2)
        except OSError:
            pass
        return []

    try:
        with open(path, "r", encoding="utf-8") as fh:
            payload = json.load(fh)
    except (ValueError, OSError):
        return []

    if not isinstance(payload, list):
        return []

    records: List[Dict[str, Any]] = []
    for item in payload:
        if isinstance(item, dict):
            records.append(prepare_record(item))
    return records


def save_records(records: Iterable[Dict[str, Any]], storage_path: Optional[str] = None) -> str:
    path = storage_path or get_data_path()
    prepared = [prepare_record(record) for record in records]
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(prepared, fh, ensure_ascii=False, indent=2)
    return path


def find_header_row(ws) -> int:
    max_row = ws.max_row
    for row_idx in range(1, max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if isinstance(cell_value, str) and "COD" in cell_value.upper():
                return row_idx
    return 1


def cell_to_text(cell_value: Any) -> str:
    if cell_value is None:
        return ""
    if isinstance(cell_value, datetime):
        return cell_value.date().strftime("%d/%m/%Y")
    if isinstance(cell_value, date):
        return cell_value.strftime("%d/%m/%Y")
    return str(cell_value)


def import_excel_file(path: str) -> List[Dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header_row = find_header_row(ws)
    headers = [cell_to_text(ws.cell(row=header_row, column=column).value) for column in range(1, ws.max_column + 1)]

    mapping = {}
    for index, header in enumerate(headers, start=1):
        normalized = header.strip().upper()
        if normalized in {field.upper() for field in EXCEL_IMPORT_FIELDS}:
            mapping[normalized] = index

    records: List[Dict[str, Any]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        row = {}
        for field_key in EXCEL_IMPORT_FIELDS:
            column_index = mapping.get(field_key.upper())
            if column_index is None:
                continue
            row[field_key] = ws.cell(row=row_idx, column=column_index).value

        if not row or all(cell_to_text(row.get(field, "")).strip() == "" for field in EXCEL_IMPORT_FIELDS):
            continue

        prepared = {
            "COD": cell_to_text(row.get("COD", "")),
            "DATA ENTRADA": cell_to_text(row.get("DATA ENTRADA", "")),
            "CLIENTES": cell_to_text(row.get("CLIENTES", "")),
            "PACIENTE": cell_to_text(row.get("PACIENTE", "")),
            "SUP/INF": cell_to_text(row.get("SUP/INF", "")),
            "SAIDA": cell_to_text(row.get("SAIDA", "")),
            "OBSERVAÇÕES": "",
            "ENTREGUE": False,
        }
        records.append(prepare_record(prepared))

    wb.close()
    return records


def export_excel_file(records: Iterable[Dict[str, Any]], output_path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "TRABALHOS"
    ws.append(EXCEL_EXPORT_FIELDS)

    for record in records:
        prepared = prepare_record(record)
        row = [
            prepared.get("COD", ""),
            prepared.get("DATA ENTRADA", ""),
            prepared.get("CLIENTES", ""),
            prepared.get("PACIENTE", ""),
            prepared.get("SUP/INF", ""),
            prepared.get("SAIDA", ""),
            prepared.get("STATUS", ""),
            prepared.get("OBSERVAÇÕES", ""),
        ]
        ws.append(row)

    wb.save(output_path)
    wb.close()
    return output_path
